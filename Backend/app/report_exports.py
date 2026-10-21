"""Immutable report files. Rendering accepts saved evidence, never SQL or model tools."""

import asyncio
import hashlib
import io
import json
import multiprocessing as mp
import os
import re
import tempfile
import textwrap
import time
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from app.chart_presentation import csv_line, present
from app.errors import AppError

FORMATS = {
    "pdf": ("application/pdf", "pdf"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "data": ("application/zip", "zip"),
}


def zip_entry(archive, name, data):
    entry = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
    entry.compress_type = zipfile.ZIP_DEFLATED
    archive.writestr(entry, data)


def format_axis(value, measure):
    unit = measure.get("unit", "number")
    if unit == "currency":
        return f"{measure.get('currency') or 'Currency'} {value:,.0f}"
    if unit == "percent":
        scaled = value * 100 if measure.get("percent_scale") == "fraction" else value
        return f"{scaled:,.1f}%"
    return f"{value:,.0f}" if unit == "count" else f"{value:,.2f}"


def chart_image(artifact, rows):
    os.environ.setdefault("MPLCONFIGDIR", "/tmp/querylane-matplotlib")
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    from matplotlib.ticker import FuncFormatter

    spec = artifact["visualization"]
    if spec["type"] in {"table", "kpi"} or not rows:
        return None
    view = spec["type"]
    presentation = present(artifact, rows, view=view)
    if view == "pie" and not presentation["composition"]["valid"]:
        return None
    if view == "scatter" and len(spec.get("y", [])) < 2:
        return None
    data = presentation["rows"]
    if not data:
        return None
    xs, ys = spec.get("x"), spec.get("y", [])
    if not xs or not ys:
        return None
    figure = Figure(figsize=(9, 4.4), layout="constrained")
    axes = figure.subplots()
    axes.spines[["top", "right"]].set_visible(False)
    axes.spines[["left", "bottom"]].set_color("#DCE3ED")
    axes.tick_params(labelsize=9, colors="#596A82")
    colors = ["#315DEE", "#176F69", "#8B5CF6", "#B45309"]
    for index, field in enumerate(ys):
        values = [float(row[field]) if row.get(field) is not None else float("nan") for row in data]
        label = next((m.get("label") for m in spec.get("measures", []) if m["field"] == field), None) or field
        if view == "bar":
            axes.bar([str(r[xs]) for r in data], values, alpha=0.75, label=label, color=colors[index % 4])
        elif view == "scatter" and len(ys) >= 2:
            axes.scatter([r[ys[0]] for r in data], [r[ys[1]] for r in data], s=16)
            break
        elif view == "pie":
            axes.pie(values, labels=[str(r[xs]) for r in data])
            break
        else:
            axes.plot(
                [str(r[xs]) for r in data],
                values,
                label=label,
                linewidth=1.8,
                linestyle="--" if field in {"forecast", "scenario", "planned"} else "-",
                color=colors[index % 4],
            )
    if view not in {"pie", "scatter"}:
        ticks = list(range(0, len(data), max(1, len(data) // 6)))
        axes.set_xticks(ticks, [str(data[i][xs]) for i in ticks], rotation=25, ha="right")
    measures = spec.get("measures", [])
    if measures and view != "pie":
        y_measure = measures[1] if view == "scatter" and len(measures) > 1 else measures[0]
        axes.yaxis.set_major_formatter(FuncFormatter(lambda value, _: format_axis(value, y_measure)))
        axes.set_ylabel(y_measure.get("label") or y_measure["field"])
        if view == "scatter" and len(measures) > 1:
            axes.xaxis.set_major_formatter(FuncFormatter(lambda value, _: format_axis(value, measures[0])))
            axes.set_xlabel(measures[0].get("label") or measures[0]["field"])
        else:
            axes.set_xlabel(xs)
    prediction = artifact.get("prediction") or {}
    if prediction and view not in {"pie", "scatter", "bar"}:
        cutoff = prediction.get("cutoff", "")
        historical = [i for i, row in enumerate(data) if str(row.get(xs, "")) <= cutoff]
        if historical and historical[-1] < len(data) - 1:
            axes.axvline(historical[-1] + 0.5, color="#596A82", linestyle=":", linewidth=1)
            axes.text(
                0.01,
                0.97,
                f"Observed through {cutoff} · dashed = predicted",
                transform=axes.transAxes,
                fontsize=8,
                color="#596A82",
                va="top",
                bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.9},
            )
    if prediction.get("interval"):
        lo = [
            float(r.get("interval_lower")) if r.get("interval_lower") is not None else float("nan")
            for r in data
        ]
        hi = [
            float(r.get("interval_upper")) if r.get("interval_upper") is not None else float("nan")
            for r in data
        ]
        axes.fill_between(range(len(data)), lo, hi, color=colors[1], alpha=0.12, label="Prediction interval")
    # Artifact titles also appear as document headings. Wrap the chart copy so
    # long model-generated titles cannot extend beyond the image boundary.
    axes.set_title(textwrap.fill(artifact["title"], width=68), loc="left", fontsize=11)
    if view != "pie":
        axes.legend(loc="best", fontsize=8)
        axes.grid(axis="y", alpha=0.15)
    buffer = io.BytesIO()
    figure.savefig(buffer, format="png", dpi=150, metadata={"Software": "Querylane"})
    figure.clear()
    return buffer.getvalue()


def chart_images(artifact, rows):
    measures = artifact["visualization"].get("measures", [])
    if artifact["visualization"]["type"] in {"scatter", "pie", "table", "kpi"}:
        image = chart_image(artifact, rows)
        return [image] if image else []
    groups = {}
    for measure in measures:
        key = (measure["unit"], measure.get("currency"), measure.get("percent_scale"))
        groups.setdefault(key, []).append(measure)
    if len(groups) <= 1:
        image = chart_image(artifact, rows)
        return [image] if image else []
    images = []
    for key, group in groups.items():
        selected = artifact | {
            "visualization": artifact["visualization"] | {"y": [m["field"] for m in group], "measures": group}
        }
        selected["title"] = artifact["title"] + " — " + " / ".join(str(k) for k in key if k)
        image = chart_image(selected, rows)
        if image:
            images.append(image)
    return images


def plain(markdown):
    return re.sub(r"(?:\*\*|__|`|^#{1,6}\s+)", "", markdown, flags=re.M)


def _render(document, artifacts):
    """Render a briefing from frozen inputs; no network, SQL or model dependencies."""
    from copy import deepcopy
    from datetime import UTC, datetime

    os.environ.setdefault("MPLCONFIGDIR", "/tmp/querylane-matplotlib")
    import matplotlib
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Image,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )
    from reportlab.platypus.tableofcontents import TableOfContents

    # Chart presentation attaches display-only details to rows. Isolate those
    # additions from captured inputs and the full-row CSV evidence.
    artifacts = deepcopy(artifacts)
    # Matplotlib already packages these fonts in the deployed image. Unicode
    # evidence is kept selectable, with no network font download or host dependency.
    font_dir = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    for face, filename in [("QL", "DejaVuSans.ttf"), ("QL-Bold", "DejaVuSans-Bold.ttf")]:
        if face not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(face, str(font_dir / filename)))
    pdfmetrics.registerFontFamily("QL", normal="QL", bold="QL-Bold", italic="QL", boldItalic="QL-Bold")
    navy, cobalt, teal = "18283B", "315DEE", "176F69"
    ink = colors.HexColor("#" + navy)
    muted = colors.HexColor("#596A82")
    rule = colors.HexColor("#DCE3ED")
    pale = colors.HexColor("#F7F9FC")
    width = A4[0] - 88
    styles = {
        "body": ParagraphStyle(
            "QLBody", fontName="QL", fontSize=9, leading=14, textColor=ink, spaceAfter=8, splitLongWords=True
        ),
        "title": ParagraphStyle(
            "QLTitle", fontName="QL-Bold", fontSize=25, leading=32, textColor=ink, spaceAfter=18
        ),
        "h1": ParagraphStyle(
            "QLSection",
            fontName="QL-Bold",
            fontSize=15,
            leading=21,
            textColor=ink,
            spaceBefore=22,
            spaceAfter=12,
            keepWithNext=True,
            borderColor=rule,
            borderWidth=0.6,
            borderPadding=8,
            backColor=pale,
        ),
        "h2": ParagraphStyle(
            "QLSubsection",
            fontName="QL-Bold",
            fontSize=11,
            leading=16,
            textColor=ink,
            spaceBefore=14,
            spaceAfter=8,
            keepWithNext=True,
        ),
        "small": ParagraphStyle(
            "QLSmall",
            fontName="QL",
            fontSize=7.5,
            leading=11,
            textColor=muted,
            spaceAfter=8,
            splitLongWords=True,
        ),
        "cell": ParagraphStyle(
            "QLCell", fontName="QL", fontSize=7.5, leading=11, textColor=ink, splitLongWords=True
        ),
        "code": ParagraphStyle(
            "QLCode",
            fontName="QL",
            fontSize=7.5,
            leading=11,
            textColor=ink,
            backColor=pale,
            borderPadding=8,
            spaceAfter=12,
            splitLongWords=True,
        ),
    }
    word = Document()
    section = word.sections[0]
    section.page_width, section.page_height = Inches(A4[0] / 72), Inches(A4[1] / 72)
    section.left_margin = section.right_margin = Inches(44 / 72)
    section.top_margin = section.bottom_margin = Inches(0.7)
    for name in ["Normal", "Body Text", "List Bullet", "List Number"]:
        style = word.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(10)
        style.font.color.rgb = RGBColor.from_string(navy)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.2
    for name, size in [("Title", 28), ("Heading 1", 17), ("Heading 2", 12)]:
        style = word.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor.from_string(navy)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(16)
        style.paragraph_format.space_after = Pt(8)
    header = section.header.paragraphs[0]
    header.text = "Querylane  /  SAVED INTELLIGENCE"
    header.style = word.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.text = f"Querylane · Version {document['version']} · Page "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    word.core_properties.title = document["title"]
    word.core_properties.author = "Querylane"
    word.core_properties.subject = f"Saved report {document['report_id']} · version {document['version']}"
    # Stable document properties and ZIP entry times aid durable render recovery.
    created = datetime.fromisoformat(document["created_at"].replace("Z", "+00:00"))
    word.core_properties.created = word.core_properties.modified = created
    story = []
    headings = []
    figure_number = 0

    def inline(value):
        value = escape(str(value))
        value = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", value)
        value = re.sub(r"`([^`]+)`", r"<font color='#176F69'>\1</font>", value)
        # Links are human-readable; no HTML or arbitrary ReportLab tags accepted.
        return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", value)

    def paragraph(text, kind="body", word_style=None):
        story.append(Paragraph(inline(text).replace("\n", "<br/>"), styles[kind]))
        wp = word.add_paragraph(style=word_style or "Normal")
        for index, part in enumerate(re.split(r"(\*\*.*?\*\*|`[^`]+`)", str(text))):
            if part.startswith("**") and part.endswith("**"):
                wp.add_run(part[2:-2]).bold = True
            else:
                wp.add_run(part.strip("`") if part.startswith("`") else part)
        if kind == "small":
            for run in wp.runs:
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor.from_string("596A82")
        if kind == "code":
            for run in wp.runs:
                run.font.name, run.font.size = "Consolas", Pt(8)
        return wp

    def heading(text, level=1):
        hp = Paragraph(escape(str(text)), styles["h1" if level == 1 else "h2"])
        if level == 1:
            hp._ql_bookmark = f"section-{len(headings)}"
            headings.append((str(text), hp._ql_bookmark))
        story.append(hp)
        word.add_heading(str(text), level)

    def table(headers, rows, numeric=()):
        if not headers:
            return
        cells = [
            [Paragraph(escape(str(value)).replace("\n", "<br/>"), styles["cell"]) for value in row]
            for row in [headers] + rows
        ]
        for row in cells[1:]:
            for col in numeric:
                if col < len(row):
                    row[col].style = ParagraphStyle("QLNumber", parent=styles["cell"], alignment=TA_RIGHT)
        pdf_table = Table(
            cells, repeatRows=1, colWidths=[width / len(headers)] * len(headers), hAlign="LEFT", splitInRow=1
        )
        pdf_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF0FF")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, pale]),
                    ("LINEBELOW", (0, 0), (-1, 0), 0.7, rule),
                    ("LINEBELOW", (0, 1), (-1, -1), 0.3, rule),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
        story.extend([pdf_table, Spacer(1, 12)])
        wt = word.add_table(rows=1, cols=len(headers))
        wt.style = "Light Shading Accent 1"
        repeat = OxmlElement("w:tblHeader")
        wt.rows[0]._tr.get_or_add_trPr().append(repeat)
        for cell, value in zip(wt.rows[0].cells, headers):
            cell.text = str(value)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for values in rows:
            row = wt.add_row()
            for index, (cell, value) in enumerate(zip(row.cells, values)):
                cell.text = str(value)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(4)
                    if index in numeric:
                        para.alignment = 2
                    for run in para.runs:
                        run.font.size = Pt(8)
        word.add_paragraph()

    def markdown(text):
        lines = str(text).splitlines()
        index = 0
        while index < len(lines):
            line = lines[index].strip()
            if not line:
                index += 1
                continue
            if line.startswith("```"):
                block = []
                index += 1
                while index < len(lines) and not lines[index].strip().startswith("```"):
                    block.append(lines[index])
                    index += 1
                paragraph("\n".join(block), "code")
            elif index + 1 < len(lines) and "|" in line and re.match(r"^[\s|:\-]+$", lines[index + 1]):
                headers = [part.strip() for part in line.strip("|").split("|")]
                index += 2
                rows = []
                while index < len(lines) and "|" in lines[index]:
                    parts = [part.strip() for part in lines[index].strip().strip("|").split("|")]
                    rows.append((parts + [""] * len(headers))[: len(headers)])
                    index += 1
                table(headers, rows)
                continue
            elif re.match(r"^#{1,6}\s+", line):
                heading(re.sub(r"^#{1,6}\s+", "", line), 2)
            elif re.match(r"^(?:[-*•]|\d+[.)])\s+", line):
                numbered = bool(re.match(r"^\d", line))
                value = re.sub(r"^(?:[-*•]|\d+[.)])\s+", "", line)
                prefix = re.match(r"^\d+[.)]", line).group() if numbered else "•"
                story.append(Paragraph(inline(value), styles["body"], bulletText=prefix))
                word.add_paragraph(plain(value), style="List Number" if numbered else "List Bullet")
            else:
                paragraph(line)
            index += 1

    def value_text(value, column, artifact):
        if value is None:
            return "Not available"
        if isinstance(value, (float, int)) and not isinstance(value, bool):
            measure = next(
                (m for m in artifact["visualization"].get("measures", []) if m["field"] == column), {}
            )
            if measure.get("unit") == "currency":
                return f"{measure.get('currency') or 'Currency'} {value:,.2f}"
            if measure.get("unit") == "percent":
                return f"{value * (100 if measure.get('percent_scale') == 'fraction' else 1):,.2f}%"
            return f"{value:,}" if isinstance(value, int) else f"{value:,.2f}"
        text = str(value)
        return text if len(text) <= 160 else text[:157] + "..."

    def prediction_evidence(prediction):
        comparison = prediction.get("comparison") or {}
        adjusted = prediction.get("kind") == "adjusted_plan"
        heading(
            "Planning assumptions — not statistical predictions"
            if adjusted
            else "Forecast method and validation",
            2,
        )
        if prediction.get("status") == "allocated":
            paragraph("Allocated estimate; this node was not independently modeled or validated.")
        table(
            ["Observed through", "Horizon", "Selected method"],
            [
                [
                    prediction.get("cutoff", "Not recorded"),
                    f"{prediction.get('horizon', '—')} {prediction.get('grain', '')} periods",
                    (
                        prediction.get("model")
                        or (
                            "TimesFM 3.0"
                            if prediction.get("method") == "timesfm_3"
                            else prediction.get("method", "Not recorded")
                        )
                    )
                    + (" · Preview" if prediction.get("preview") else ""),
                ]
            ],
        )
        paragraph(prediction.get("selection_reason") or "Selection rationale not recorded.")
        paragraph(
            (
                "Solid lines show observations; dashed lines show user planning assumptions. "
                if adjusted
                else "Solid lines show observations; dashed lines show predictions. "
            )
            + (
                f"Saved interval level: {prediction['interval'].get('level')}."
                if prediction.get("interval")
                else "Validated intervals unavailable; uncertainty is not quantified."
            ),
            "small",
        )
        if comparison.get("reason"):
            paragraph(comparison["reason"], "small")
        if comparison.get("scope_note"):
            paragraph(comparison["scope_note"], "small")
        scores = prediction.get("scores", {})
        if scores:
            paragraph(
                "Chronological backtests: errors use the saved target units; lower is better, not an accuracy percentage.",
                "small",
            )
            table(
                ["Candidate", "MAE", "RMSE", "Outcome"],
                [
                    [
                        name,
                        f"{score['mae']:,.2f}",
                        f"{score['rmse']:,.2f}",
                        "Selected" if name == prediction.get("method") else "Compared",
                    ]
                    for name, score in scores.items()
                ],
                numeric=(1, 2),
            )
            if document.get("style") == "technical":
                for name, score in scores.items():
                    if score.get("folds"):
                        heading(f"{name} · held-out folds", 2)
                        table(
                            ["Training end", "Test start", "Test end", "MAE"],
                            [
                                [
                                    fold.get("train_end", "—"),
                                    fold.get("test_start", "—"),
                                    fold.get("test_end", "—"),
                                    f"{fold['mae']:,.2f}",
                                ]
                                for fold in score["folds"]
                            ],
                            numeric=(3,),
                        )
        if prediction.get("kind") == "scenario":
            heading("Conditional scenario assumptions", 2)
            paragraph(
                "Baseline and alternative paths are explicit assumptions, not causal lift estimates. Both paths use the same saved scenario method and history."
            )
            for assumption in prediction.get("assumptions", []):
                paragraph(assumption)
            baseline, alternative = (
                prediction.get("future_drivers", []),
                prediction.get("scenario_drivers", []),
            )
            drivers = prediction.get("drivers", [])
            if drivers:
                table(
                    ["Period", "Driver", "Baseline", "Alternative"],
                    [
                        [
                            i + 1,
                            driver,
                            path.get(driver, "Not available"),
                            alternative[i].get(driver, "Not available")
                            if i < len(alternative)
                            else "Not available",
                        ]
                        for i, path in enumerate(baseline)
                        for driver in drivers
                    ],
                    numeric=(2, 3),
                )
        for limitation in prediction.get("limitations", []):
            paragraph(limitation, "small")
        if document.get("style") == "technical":
            comparison = prediction.get("comparison") or {}
            table(
                ["Provider / model", "Context", "Input hash"],
                [
                    [
                        f"{prediction.get('provider', 'Not recorded')} / {prediction.get('model') or prediction.get('method')}",
                        str(comparison.get("context", "Not recorded")),
                        comparison.get("input_hash") or prediction.get("data_hash", "Not recorded"),
                    ]
                ],
            )
            paragraph(
                f"Selection policy: {comparison.get('selection_policy_version', 'Not recorded')}. Comparison status: {comparison.get('status', 'Not recorded')}.",
                "small",
            )

    paragraph("QUERYLANE  /  SAVED INTELLIGENCE", "small")
    story.append(Paragraph(escape(document["title"]), styles["title"]))
    word.add_heading(document["title"], 0)
    table(
        ["Report style", "Saved version", "Created"],
        [
            [
                document.get("style", "business").title(),
                document["version"],
                created.astimezone(UTC).strftime("%d %b %Y · %H:%M UTC")
                if created.tzinfo
                else created.strftime("%d %b %Y"),
            ]
        ],
    )
    paragraph(
        f"Public report · Version {document['version']} · {document['data_profile']} data profile · ID {document['report_id']}",
        "small",
    )
    if document.get("data_profile") == "timeseries":
        paragraph(
            "Synthetic retail data · results support demonstration and evaluation, not guaranteed business outcomes.",
            "small",
        )
    heading("Executive summary")
    markdown(document["summary"])
    if document.get("claims"):
        heading("Saved verified highlights", 2)
        table(
            ["Finding", "Saved value"],
            [
                [claim["label"], claim["formatted"]]
                for claim in [
                    claim
                    for claim in document["claims"]
                    if not re.search(r"number of saved result rows|dimension labels", claim["label"], re.I)
                ][:3]
            ],
        )
    if document.get("scope"):
        heading("Analysis scope", 2)
        markdown(document["scope"])
    if len(document["sections"]) >= 4:
        story.append(PageBreak())
        paragraph("CONTENTS", "small")
        toc = TableOfContents()
        toc.levelStyles = [ParagraphStyle("QLContents", parent=styles["body"], spaceBefore=9)]
        story.append(toc)
        word.add_page_break()
        word.add_heading("Contents", 1)
        # Word contains usable internal hyperlinks immediately, plus standard
        # headings for its navigation pane. No field update is required to read it.
        toc_paragraphs = []
        for index, title in enumerate(
            [s["title"] for s in document["sections"]] + ["Evidence and methodology"]
        ):
            wp = word.add_paragraph()
            link = OxmlElement("w:hyperlink")
            link.set(qn("w:anchor"), f"ql-section-{index}")
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), cobalt)
            props.append(color)
            run.append(props)
            text = OxmlElement("w:t")
            text.text = title
            run.append(text)
            link.append(run)
            wp._p.append(link)
            toc_paragraphs.append(wp)
        story.append(PageBreak())
    for index, item in enumerate(document["sections"]):
        heading(f"{index + 1:02d}  {item['title']}")
        wp = word.paragraphs[-1]
        if index == 0 and len(document["sections"]) >= 4:
            wp.paragraph_format.page_break_before = True
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), str(index))
        start.set(qn("w:name"), f"ql-section-{index}")
        end.set(qn("w:id"), str(index))
        wp._p.insert(1 if wp._p.pPr is not None else 0, start)
        wp._p.append(end)
        if item.get("kind") == "recommendation":
            paragraph("Recommendation · analyst interpretation", "small")
        markdown(item["markdown"])
        for artifact_id in item.get("artifact_ids", []):
            artifact, rows = artifacts[artifact_id]
            heading(artifact["title"], 2)
            paragraph(
                f"Saved result · {len(rows):,} rows · {'Complete inputs' if artifact['complete'] else 'PARTIAL — not a population total'}",
                "small",
            )
            if artifact.get("prediction"):
                prediction_evidence(artifact["prediction"])
            if artifact["visualization"]["type"] not in {"table", "kpi"} and rows:
                coverage = present(artifact, rows, view=artifact["visualization"]["type"])["coverage"]
                paragraph(
                    f"Chart coverage: {coverage['plotted_rows']} plotted of {coverage['selected_rows']} saved observations; {coverage['omitted_categories']} omitted categories. "
                    + (
                        "Extrema-preserving sampling applied."
                        if coverage["sampled"]
                        else "No visual sampling."
                    ),
                    "small",
                )
            images = chart_images(artifact, rows)
            for image in images:
                figure_number += 1
                story.append(Image(io.BytesIO(image), width=width, height=width * 4.4 / 9))
                word.add_picture(io.BytesIO(image), width=Inches(width / 72))
                paragraph(f"Figure {figure_number} · {artifact['title']} · saved evidence", "small")
            if not images and artifact["visualization"]["type"] not in {"table", "kpi"}:
                paragraph(
                    "This saved scope cannot be represented faithfully by the requested static chart. Inspect its table and full CSV instead.",
                    "small",
                )
            prediction = artifact.get("prediction") or {}
            dimension = artifact["visualization"].get("x")
            projected_fields = [
                name
                for name in ("forecast", "scenario", "planned")
                if any(row.get(name) is not None for row in rows)
            ]
            projected = (
                [
                    row
                    for row in rows
                    if dimension
                    and str(row.get(dimension, "")) > prediction.get("cutoff", "")
                    and any(row.get(name) is not None for name in projected_fields)
                ]
                if prediction
                else []
            )
            if projected and projected_fields:
                heading("Saved future point estimates", 2)
                future_columns = [dimension] + projected_fields
                table(
                    [name.replace("_", " ").title() for name in future_columns],
                    [
                        [value_text(row.get(name), name, artifact) for name in future_columns]
                        for row in projected
                    ],
                    numeric=tuple(range(1, len(future_columns))),
                )
                paragraph(
                    "These are saved predictions or planning assumptions after the observation cutoff, not observed outcomes. No interval is implied by point estimates.",
                    "small",
                )
            names = [column["name"] for column in artifact["columns"]][:8]
            table(
                [name.replace("_", " ").title() for name in names],
                [[value_text(row.get(name), name, artifact) for name in names] for row in rows[:20]],
                numeric=tuple(
                    i
                    for i, name in enumerate(names)
                    if any(isinstance(row.get(name), (int, float)) for row in rows[:20])
                ),
            )
            paragraph(
                "Document preview: up to 20 rows and 8 columns; text cells over 160 characters are abbreviated. Full saved rows and original precision are in the data ZIP. "
                + f"Evidence ID: {artifact_id}",
                "small",
            )
    heading("Evidence and methodology")
    # Anchor the final contents entry without depending on Word field refresh.
    wp = word.paragraphs[-1]
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(len(document["sections"])))
    start.set(qn("w:name"), f"ql-section-{len(document['sections'])}")
    end.set(qn("w:id"), str(len(document["sections"])))
    wp._p.insert(1 if wp._p.pPr is not None else 0, start)
    wp._p.append(end)
    for title, items in [
        ("Definitions", document.get("definitions", [])),
        ("Assumptions", document.get("assumptions", [])),
        ("Limitations", document["limitations"]),
    ]:
        heading(title, 2)
        for item in items:
            paragraph(item)
    heading("Verified calculations", 2)
    for claim in document.get("claims", []):
        paragraph(f"{claim['label']}: {claim['formatted']} — {claim['calculation']}")
    paragraph(
        f"Data profile: {document['data_profile']}. Model: {document['model']}. {document['mode_note']}",
        "small",
    )
    for artifact_id, (artifact, _) in artifacts.items():
        heading(f"Evidence: {artifact['title']}", 2)
        paragraph(f"Artifact {artifact_id}", "small")
        for source in artifact["provenance"]:
            paragraph(
                f"{source['source']} · {source['executed_at']} · {source.get('job_id') or 'read-only query'}",
                "small",
            )
            paragraph(source["sql"], "code")
        for assumption in artifact.get("assumptions", []):
            paragraph(assumption)
    paragraph(
        "Source timestamps belong to independent saved executions; they do not imply one synchronized snapshot.",
        "small",
    )
    for skill in document.get("skills_used", []):
        paragraph(f"Skill: {skill['name']} · revision {skill['revision']}", "small")

    class BriefingTemplate(SimpleDocTemplate):
        def afterFlowable(self, flowable):
            key = getattr(flowable, "_ql_bookmark", None)
            if key:
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(flowable.getPlainText(), key, level=0)
                self.notify("TOCEntry", (0, flowable.getPlainText(), self.page, key))

    def page_furniture(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(rule)
        canvas.line(44, A4[1] - 31, A4[0] - 44, A4[1] - 31)
        canvas.setFont("QL-Bold", 8)
        canvas.setFillColor(colors.HexColor("#" + teal))
        canvas.drawString(44, A4[1] - 23, "QUERYLANE / SAVED INTELLIGENCE")
        canvas.setFont("QL", 7)
        canvas.setFillColor(muted)
        canvas.drawString(44, 25, f"Public report · v{document['version']}")
        canvas.drawRightString(A4[0] - 44, 25, f"Page {doc.page}")
        canvas.restoreState()

    pdf_buffer = io.BytesIO()
    pdf = BriefingTemplate(
        pdf_buffer,
        pagesize=A4,
        title=document["title"],
        author="Querylane",
        subject=f"Report {document['report_id']} version {document['version']}",
        invariant=1,
        leftMargin=44,
        rightMargin=44,
        topMargin=48,
        bottomMargin=44,
    )
    pdf.multiBuild(story, onFirstPage=page_furniture, onLaterPages=page_furniture)
    wb = io.BytesIO()
    word.save(wb)
    normalized = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(wb.getvalue())) as source,
        zipfile.ZipFile(normalized, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in source.namelist():
            zip_entry(target, name, source.read(name))
    data = io.BytesIO()
    manifest = {
        "report_id": document["report_id"],
        "version": document["version"],
        "created_at": document["created_at"],
        "data_profile": document["data_profile"],
        "claims": document["claims"],
        "presentation_version": 2,
        "datasets": [],
    }
    dictionary = [["file", "column", "type", "unit", "currency", "percent_scale"]]
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as archive:
        for artifact_id, (artifact, rows) in artifacts.items():
            columns = [column["name"] for column in artifact["columns"]]
            csv = (
                "\ufeff"
                + csv_line(columns)
                + "".join(csv_line([row.get(column) for column in columns]) for row in rows)
            )
            filename = f"{artifact_id}.csv"
            zip_entry(archive, filename, csv)
            manifest["datasets"].append(
                {
                    "file": filename,
                    "title": artifact["title"],
                    "rows": len(rows),
                    "sha256": hashlib.sha256(csv.encode()).hexdigest(),
                    "complete": artifact["complete"],
                    "provenance": artifact["provenance"],
                    "prediction": artifact.get("prediction"),
                }
            )
            for column in artifact["columns"]:
                measure = next(
                    (
                        m
                        for m in artifact["visualization"].get("measures", [])
                        if m["field"] == column["name"]
                    ),
                    {},
                )
                dictionary.append(
                    [
                        filename,
                        column["name"],
                        column.get("type", "not recorded"),
                        measure.get("unit", "not recorded"),
                        measure.get("currency", ""),
                        measure.get("percent_scale", ""),
                    ]
                )
        zip_entry(archive, "manifest.json", json.dumps(manifest, indent=2))
        zip_entry(archive, "report.json", json.dumps(document, indent=2, ensure_ascii=True))
        zip_entry(archive, "data_dictionary.csv", "\ufeff" + "".join(csv_line(row) for row in dictionary))
        index = "\n".join(
            f"- {item['file']}: {item['title']} ({item['rows']} rows; {'complete' if item['complete'] else 'PARTIAL'})"
            for item in manifest["datasets"]
        )
        zip_entry(
            archive,
            "README.md",
            f"# {document['title']}\n\nQuerylane public report · version {document['version']}\n\n{index}\n\nCSV files preserve all saved rows, original columns and numeric precision. Preview formatting in PDF/Word does not change CSV evidence.\n\nmanifest.json maps evidence to files, source timestamps, completeness, SHA-256 and forecast metadata. data_dictionary.csv lists recorded column types and units; unrecorded units are not guessed. report.json contains the saved narrative, definitions and assumptions.\n\nForecasts and scenarios are predictions, not observations or causal effects. Consult saved cutoff, assumptions and interval limitations. No download performs analysis or inference.\n",
        )
    return {"pdf": pdf_buffer.getvalue(), "docx": normalized.getvalue(), "data": data.getvalue()}


def _render_entry(pipe, document, artifacts):
    try:
        pipe.send({"files": _render(document, artifacts)})
    except Exception:
        pipe.send({"error": "Document rendering failed. Saved evidence is preserved for recovery."})
    finally:
        pipe.close()


async def render(document, artifacts, timeout=60):
    context = mp.get_context("spawn")
    parent, child = context.Pipe(duplex=False)
    process = context.Process(target=_render_entry, args=(child, document, artifacts), daemon=True)
    process.start()
    child.close()
    started = time.monotonic()
    try:
        while not parent.poll():
            if time.monotonic() - started >= timeout or not process.is_alive():
                raise AppError(
                    "Report export deadline reached. Evidence is preserved.", 503, "report_render_error"
                )
            await asyncio.sleep(0.05)
        value = parent.recv()
        if "error" in value:
            raise AppError(value["error"], 503, "report_render_error")
        return value["files"]
    finally:
        if process.is_alive():
            process.terminate()
        await asyncio.to_thread(process.join, 2)
        if process.is_alive():
            process.kill()
            await asyncio.to_thread(process.join, 2)
        parent.close()


class ExportStorage:
    def __init__(self, settings, credentials=None):
        self.settings, self.credentials = settings, credentials
        self.client = None
        if settings.report_bucket:
            from google.cloud import storage

            self.client = storage.Client(project=settings.google_cloud_project, credentials=credentials)

    def _path(self, key):
        root = Path(self.settings.report_local_directory).resolve()
        path = (root / key).resolve()
        if not path.is_relative_to(root):
            raise AppError("Invalid export path.", 422)
        return path

    async def save(self, key, content, content_type):
        sha = hashlib.sha256(content).hexdigest()

        def write():
            if self.client:
                from google.api_core.exceptions import PreconditionFailed

                blob = self.client.bucket(self.settings.report_bucket).blob(key)
                blob.metadata = {"sha256": sha}
                try:
                    blob.upload_from_string(
                        content, content_type=content_type, if_generation_match=0, timeout=30
                    )
                except PreconditionFailed:
                    blob.reload(timeout=15)
                    if blob.metadata.get("sha256") != sha:
                        raise AppError(
                            "An immutable export already exists with different content.",
                            409,
                            "report_immutable",
                        )
            else:
                path = self._path(key)
                path.parent.mkdir(parents=True, exist_ok=True)
                handle, temporary = tempfile.mkstemp(prefix=".report-", dir=path.parent)
                try:
                    with os.fdopen(handle, "wb") as file:
                        file.write(content)
                        file.flush()
                        os.fsync(file.fileno())
                    try:
                        os.link(temporary, path)
                    except FileExistsError:
                        if hashlib.sha256(path.read_bytes()).hexdigest() != sha:
                            raise AppError("Immutable export conflict.", 409, "report_immutable")
                finally:
                    Path(temporary).unlink(missing_ok=True)

        await asyncio.to_thread(write)
        return {"key": key, "sha256": sha, "size": len(content), "content_type": content_type}

    async def read(self, metadata):
        def read():
            return (
                self.client.bucket(self.settings.report_bucket)
                .blob(metadata["key"])
                .download_as_bytes(timeout=30)
                if self.client
                else self._path(metadata["key"]).read_bytes()
            )

        try:
            value = await asyncio.to_thread(read)
        except Exception as exc:
            raise AppError(
                "This export is unavailable. Contact the operator; no queries were rerun.",
                503,
                "report_export_unavailable",
            ) from exc
        if hashlib.sha256(value).hexdigest() != metadata["sha256"]:
            raise AppError("Export integrity check failed.", 503, "report_export_integrity")
        return value

    async def delete(self, metadata):
        def remove():
            if self.client:
                from google.api_core.exceptions import NotFound

                try:
                    self.client.bucket(self.settings.report_bucket).blob(metadata["key"]).delete(timeout=30)
                except NotFound:
                    pass
            else:
                self._path(metadata["key"]).unlink(missing_ok=True)

        await asyncio.to_thread(remove)

    async def close(self):
        if self.client:
            await asyncio.to_thread(self.client.close)
